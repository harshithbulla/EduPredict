"""Create the AAT presentation PPTX for EduPredict."""

from pptx import Presentation
from pptx.util import Inches, Pt

BASE = r"C:\Users\harsh\OneDrive\Desktop\STUDENT_ML\EduPredict"
FIG = BASE + r"\results\figures"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

ACCENT = (31, 78, 121)


def add_title_box(slide, text, top=0.6, size=36, bold=True, color=ACCENT):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(1))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = __import__("pptx").dml.color.RGBColor(*color)


def add_bullets(slide, items, top=1.6, size=20):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.8), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(size)
        p.space_after = Pt(8)


def add_picture(slide, path, left, top, width, height):
    try:
        slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))
    except Exception as e:
        print("Missing image:", path, e)


# Slide 1 - Title
s = prs.slides.add_slide(BLANK)
add_title_box(s, "EduPredict", top=1.2, size=48)
add_bullets(
    s,
    [
        "Student Performance Prediction Using Linear Regression",
        "",
        "Name: ______________________",
        "USN: ______________________",
        "Department: ______________________",
        "College: ______________________",
        "Academic Year: ______________________",
    ],
    top=2.4,
    size=20,
)

# Slide 2 - Introduction
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Introduction")
add_bullets(s, [
    "Machine Learning: computers learn patterns from data.",
    "Supervised learning uses data where answers are known.",
    "Regression predicts continuous numeric values.",
    "Goal: predict final exam marks from academic behavior.",
])

# Slide 3 - Problem Statement
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Problem Statement")
add_bullets(s, [
    "Predict a student's final exam marks using:",
    "   Study hours per day",
    "   Attendance percentage",
    "   Previous / internal marks",
    "Use Linear Regression (the syllabus algorithm) from scikit-learn.",
])

# Slide 4 - Objectives
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Objectives")
add_bullets(s, [
    "Prepare a realistic dataset of 200 students.",
    "Explore and clean the data.",
    "Train a Linear Regression model.",
    "Predict final marks for unseen test data.",
    "Evaluate using MAE, MSE, RMSE and R2.",
])

# Slide 5 - Dataset & Features
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Dataset & Features")
add_bullets(s, [
    "200 records, 4 columns, all numeric.",
    "Study_Hours: approx. 1 - 10 hours/day",
    "Attendance_Percentage: approx. 50 - 100%",
    "Previous_Marks: approx. 30 - 95",
    "Final_Marks (target): approx. 30 - 100",
    "Split: 160 training / 40 testing, random_state = 42",
])

# Slide 6 - Methodology
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Methodology / Workflow")
add_bullets(s, [
    "Dataset -> Data Preprocessing",
    "-> Exploratory Data Analysis",
    "-> Feature Selection -> Train-Test Split (80/20)",
    "-> Linear Regression Training -> Prediction",
    "-> Evaluation (MAE, MSE, RMSE, R2)",
    "-> Visualization -> Conclusion",
])

# Slide 7 - Algorithm
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Linear Regression Algorithm")
add_bullets(s, [
    "Fits: Final_Marks = b0 + b1*Study_Hours + b2*Attendance + b3*Previous_Marks",
    "b0 = -0.3966 (intercept)",
    "b1 = 1.1695   b2 = 0.3000   b3 = 0.4084",
    "Learned using the least-squares method in scikit-learn.",
])

# Slide 8 - EDA
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Exploratory Data Analysis")
add_picture(s, FIG + r"\correlation_heatmap.png", 0.6, 1.6, 5.8, 5.2)
add_picture(s, FIG + r"\previous_marks_vs_final_marks.png", 6.8, 1.6, 5.8, 5.2)

# Slide 9 - Results
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Model Results")
add_bullets(s, [
    "MAE  = 5.6561",
    "MSE  = 47.1436",
    "RMSE = 6.8661",
    "R2 Score = 0.6206",
    "",
    "The model explains about 62% of the variation in final marks.",
])

# Slide 10 - Actual vs Predicted
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Actual vs Predicted")
add_picture(s, FIG + r"\actual_vs_predicted.png", 3.2, 1.6, 6.5, 5.4)

# Slide 11 - Conclusion
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Conclusion & Future Scope")
add_bullets(s, [
    "Linear Regression gave a simple, interpretable predictor.",
    "All coefficients positive - matches common sense.",
    "Not perfect: average error about 5-6 marks.",
    "Future: larger real dataset, more features (assignments, sleep, consistency), comparison with other syllabus algorithms.",
])

# Slide 12 - Thank You
s = prs.slides.add_slide(BLANK)
add_title_box(s, "Thank You", top=2.8, size=44)

import os
os.makedirs(BASE + r"\presentation", exist_ok=True)
out = BASE + r"\presentation\EduPredict_AAT_Presentation.pptx"
prs.save(out)
print("Saved:", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
